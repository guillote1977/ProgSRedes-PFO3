import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTA IMPORTANTE:"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EDF2F7")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (accent color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"📌 {title} ")
    run_t.bold = True
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    run_b = p.add_run(text)
    run_b.font.name = 'Calibri'
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_presentation_docx():
    doc = Document()
    
    # Page setup: Margins 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header and Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("PFO 3 - Programación sobre Redes | IFTS N° 29")
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # -------------------------------------------------------------
    # PORTADA / CARÁTULA
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("INSTITUTO DE FORMACIÓN TÉCNICA SUPERIOR N° 29 (IFTS 29)\nTECNICATURA SUPERIOR EN DESARROLLO DE SOFTWARE")
    r_inst.font.name = 'Calibri'
    r_inst.font.size = Pt(11)
    r_inst.bold = True
    r_inst.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    p_inst.paragraph_format.space_after = Pt(36)

    p_materia = doc.add_paragraph()
    p_materia.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mat = p_materia.add_run("PROGRAMACIÓN SOBRE REDES")
    r_mat.font.name = 'Calibri'
    r_mat.font.size = Pt(13)
    r_mat.bold = True
    r_mat.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    p_materia.paragraph_format.space_after = Pt(12)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Práctica Formativa Obligatoria N° 3 (PFO 3)")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    p_title.paragraph_format.space_after = Pt(8)

    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_subtitle.add_run("Rediseño como Sistema Distribuido\n(Cliente-Servidor Concurrente con Sockets TCP y Worker Pool)")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    p_subtitle.paragraph_format.space_after = Pt(40)

    # Info table
    info_table = doc.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    col_widths = [Inches(2.2), Inches(4.3)]
    datos = [
        ("Docentes:", "Prof. Alan Portillo / Prof. Germán Ríos"),
        ("Alumno:", "Guillermo Sciulli"),
        ("Materia / Carrera:", "Programación sobre Redes / TS en Desarrollo de Software"),
        ("Fecha de Presentación:", "Octubre 2026"),
        ("Repositorio GitHub:", "https://github.com/guillote1977/ProgSRedes-PFO3"),
        ("Estado de Entrega:", "Completo (Código funcional + Diagrama + Documentación)")
    ]
    
    for row_idx, (label, val) in enumerate(datos):
        row = info_table.rows[row_idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        
        # Zebra striping
        bg_color = "F7FAFC" if row_idx % 2 == 0 else "FFFFFF"
        set_cell_background(c0, bg_color)
        set_cell_background(c1, bg_color)
        set_cell_margins(c0, top=100, bottom=100, left=120, right=120)
        set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(10.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(10.5)
        if "github.com" in val:
            r1.bold = True
            r1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            r1.underline = True
        else:
            r1.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECCIÓN 1: INTRODUCCIÓN Y ENLACE DE ENTREGA
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Enlace al Repositorio y Datos de Entrega", level=1)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(8)
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.add_run(
        "El presente documento formaliza la entrega de la Práctica Formativa Obligatoria N° 3 (PFO 3). "
        "Todo el código fuente, la implementación en Python, el diagrama de arquitectura en formatos raster y vectorial, "
        "y los ejemplos de prueba se encuentran versionados en el siguiente repositorio oficial de GitHub:"
    )
    
    add_callout(
        doc,
        "https://github.com/guillote1977/ProgSRedes-PFO3\n"
        "El repositorio es público, permitiendo la clonación directa, inspección del historial de commits y verificación de cada archivo.",
        "ENLACE OFICIAL AL REPOSITORIO GITHUB:"
    )

    # -------------------------------------------------------------
    # SECCIÓN 2: OBJETIVOS Y ALCANCE DE LA PFO 3
    # -------------------------------------------------------------
    h2 = doc.add_heading("2. Objetivos y Consignas del Trabajo", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(8)

    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.line_spacing = 1.15
    p_obj.paragraph_format.space_after = Pt(6)
    r_obj_lead = p_obj.add_run("Objetivo Principal: ")
    r_obj_lead.bold = True
    p_obj.add_run(
        "Transformar un sistema tradicional en una arquitectura distribuida escalable y desacoplada utilizando sockets "
        "y componentes modernos de infraestructura para alta disponibilidad."
    )

    p_cons = doc.add_paragraph()
    p_cons.paragraph_format.space_after = Pt(4)
    r_cons_lead = p_cons.add_run("Consignas requeridas y cumplimentadas:")
    r_cons_lead.bold = True

    items = [
        ("Diseño de Arquitectura Distribuida: ", "Diagrama integral que modela clientes móviles/web, balanceador de carga (Nginx/HAProxy), servidores workers con pool de hilos, cola de mensajes (RabbitMQ) y almacenamiento distribuido (PostgreSQL y S3)."),
        ("Implementación del Servidor Socket en Python: ", "Servidor multi-hilo no bloqueante que atiende clientes en hilos independientes y delega el procesamiento de cómputo a un ThreadPoolExecutor (Worker Pool)."),
        ("Implementación del Cliente Socket en Python: ", "Cliente interactivo con menú en consola para envío de tareas síncronas/asíncronas, consulta diferida de estado y ráfagas de prueba de carga concurrente."),
        ("Entregables: ", "Diagrama de arquitectura del sistema distribuido y código fuente completo en repositorio GitHub.")
    ]
    for bullet_title, bullet_desc in items:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(3)
        r_bt = p_item.add_run(bullet_title)
        r_bt.bold = True
        r_bt.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        p_item.add_run(bullet_desc)

    # -------------------------------------------------------------
    # SECCIÓN 3: DIAGRAMA DE ARQUITECTURA DISTRIBUIDA
    # -------------------------------------------------------------
    h3 = doc.add_heading("3. Diagrama de Arquitectura del Sistema Distribuido", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(8)

    p_diag_desc = doc.add_paragraph()
    p_diag_desc.paragraph_format.line_spacing = 1.15
    p_diag_desc.paragraph_format.space_after = Pt(8)
    p_diag_desc.add_run(
        "Se diseñó una topología distribuida de cinco niveles que desacopla la recepción de conexiones, "
        "el enrutamiento de red, el cómputo en paralelo y la persistencia de datos:"
    )

    # Insert image if exists
    img_path = "diagrama_arquitectura.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_after = Pt(12)
        r_cap = p_caption.add_run("Figura 1: Topología de 5 Capas del Sistema Distribuido (PFO 3)")
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # Component table
    comp_table = doc.add_table(rows=6, cols=3)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_table.autofit = False
    
    col_w = [Inches(1.8), Inches(1.8), Inches(2.9)]
    headers = ["Capa / Componente", "Tecnología Adoptada", "Función en la Arquitectura"]
    
    # Header row
    hdr_row = comp_table.rows[0]
    for idx, text in enumerate(headers):
        c = hdr_row.cells[idx]
        c.width = col_w[idx]
        set_cell_background(c, "1A365D")
        set_cell_margins(c, top=120, bottom=120, left=120, right=120)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    comp_data = [
        ("1. Clientes Heterogéneos", "Aplicaciones Móviles y Web", "Generan solicitudes mediante conexiones TCP directas o peticiones REST/WebSocket hacia la plataforma."),
        ("2. Balanceador de Carga", "Nginx / HAProxy", "Punto único de ingreso público. Distribuye las cargas (Round-Robin / Least Connections) y previene la caída total del sistema (anti-SPOF)."),
        ("3. Servidores Workers", "Python + ThreadPoolExecutor", "Nodos de cómputo distribuidos. Cada servidor maneja su pool de hilos para ejecutar tareas en paralelo sin bloquear la recepción."),
        ("4. Cola de Mensajes", "RabbitMQ (AMQP)", "Middleware asíncrono que desacopla la recepción del procesamiento y amortigua picos extremos de tráfico sin pérdida de datos."),
        ("5. Almacenamiento Distribuido", "PostgreSQL + AWS S3 / MinIO", "Persistencia híbrida: base de datos relacional ACID para estados y transacciones; almacenamiento de objetos para archivos pesados.")
    ]

    for row_idx, row_content in enumerate(comp_data, start=1):
        row = comp_table.rows[row_idx]
        bg_col = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_content):
            c = row.cells[col_idx]
            c.width = col_w[col_idx]
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=90, bottom=90, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            else:
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECCIÓN 4: IMPLEMENTACIÓN EN PYTHON (SOCKETS TCP)
    # -------------------------------------------------------------
    h4 = doc.add_heading("4. Implementación del Software (Sockets TCP en Python)", level=1)
    h4.paragraph_format.space_before = Pt(6)
    h4.paragraph_format.space_after = Pt(8)

    p_soft = doc.add_paragraph()
    p_soft.paragraph_format.line_spacing = 1.15
    p_soft.paragraph_format.space_after = Pt(6)
    p_soft.add_run(
        "La solución de software desarrollada cumple estrictamente con el enfoque cliente-servidor distribuido "
        "utilizando la biblioteca estándar de Python sin requerir dependencias externas de terceros."
    )

    doc.add_heading("A. Servidor Concurrente con Worker Pool (`servidor.py`)", level=2)
    p_srv = doc.add_paragraph()
    p_srv.paragraph_format.line_spacing = 1.15
    p_srv.paragraph_format.space_after = Pt(4)
    p_srv.add_run(
        "El servidor (`servidor.py`) implementa un diseño de doble concurrencia:\n"
        "1. Concurrencia a nivel de red: Cada cliente conectado se delega a un hilo daemon (`threading.Thread`) exclusivo, "
        "garantizando que múltiples clientes interactúen de forma simultánea.\n"
        "2. Concurrencia a nivel de cómputo: Las tareas solicitadas se despachan a un `ThreadPoolExecutor(max_workers=4)`, "
        "emulando el cluster de workers del diagrama arquitectónico.\n"
        "3. Almacenamiento thread-safe: Se utiliza un lock (`threading.Lock`) para proteger el diccionario global de estados de tareas, "
        "asignando a cada solicitud un identificador único con timestamp (`TASK-<milisegundos>`)."
    )

    doc.add_heading("B. Protocolo de Mensajería JSON sobre TCP", level=2)
    p_proto = doc.add_paragraph()
    p_proto.paragraph_format.line_spacing = 1.15
    p_proto.paragraph_format.space_after = Pt(4)
    p_proto.add_run(
        "La comunicación se estructura mediante tramas serializadas en formato JSON estructurado:"
    )

    # JSON table
    proto_table = doc.add_table(rows=4, cols=3)
    proto_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    proto_table.autofit = False
    
    col_w_proto = [Inches(1.8), Inches(2.2), Inches(2.5)]
    proto_headers = ["Comando", "Payload Enviado por Cliente", "Respuesta del Servidor"]
    
    for idx, text in enumerate(proto_headers):
        c = proto_table.rows[0].cells[idx]
        c.width = col_w_proto[idx]
        set_cell_background(c, "2B6CB0")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    proto_rows = [
        ("ENVIAR_TAREA", '{"tipo": "calcular_cuadrado", "payload": {"numero": 12}}', '{"status": "OK", "task_id": "TASK-1793...", "mensaje": "Tarea recibida"}'),
        ("CONSULTAR_TAREA", '{"task_id": "TASK-1793..."}', '{"status": "OK", "info": {"estado": "COMPLETADA", "resultado": 144}}'),
        ("SALIR", '{"comando": "SALIR"}', '{"status": "OK", "mensaje": "Conexión finalizada"}')
    ]

    for row_idx, row_content in enumerate(proto_rows, start=1):
        row = proto_table.rows[row_idx]
        bg_col = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_content):
            c = row.cells[col_idx]
            c.width = col_w_proto[col_idx]
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Courier New' if col_idx > 0 else 'Calibri'
            r.font.size = Pt(8.5 if col_idx > 0 else 9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            else:
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_heading("C. Cliente Interactivo y Pruebas de Carga (`cliente.py`)", level=2)
    p_cli = doc.add_paragraph()
    p_cli.paragraph_format.line_spacing = 1.15
    p_cli.paragraph_format.space_after = Pt(4)
    p_cli.add_run(
        "El cliente (`cliente.py`) provee un menú interactivo en consola que permite:\n"
        "• Enviar tareas matemáticas (`calcular_cuadrado`) y de procesamiento de cadenas (`procesar_texto`).\n"
        "• Consultar en cualquier momento el estado y resultado de una tarea mediante su ID único.\n"
        "• Ejecutar una prueba de carga concurrente enviando 3 tareas simultáneas para verificar el despacho dinámico "
        "en los 4 hilos del Worker Pool."
    )

    # -------------------------------------------------------------
    # SECCIÓN 5: GUÍA DE EJECUCIÓN Y VALIDACIÓN
    # -------------------------------------------------------------
    h5 = doc.add_heading("5. Guía de Ejecución y Validación de Pruebas", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(8)

    p_run_steps = doc.add_paragraph()
    p_run_steps.paragraph_format.line_spacing = 1.15
    p_run_steps.paragraph_format.space_after = Pt(4)
    p_run_steps.add_run(
        "Para ejecutar y probar la solución en cualquier entorno con Python 3 instalado:"
    )

    steps = [
        ("Paso 1: Clonar el Repositorio de GitHub", "git clone https://github.com/guillote1977/ProgSRedes-PFO3.git\ncd ProgSRedes-PFO3"),
        ("Paso 2: Iniciar el Servidor Distribuido", "python3 servidor.py\n(Escuchando en TCP 127.0.0.1:5000 con 4 workers activos)"),
        ("Paso 3: Iniciar el Cliente en otra terminal", "python3 cliente.py\n(Aparecerá el menú interactivo para despachar tareas y consultar estados)")
    ]

    for title, cmd in steps:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(4)
        p_st.paragraph_format.space_after = Pt(2)
        r_st = p_st.add_run(title)
        r_st.bold = True
        r_st.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
        # Code block table
        t_code = doc.add_table(rows=1, cols=1)
        t_code.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_code.autofit = False
        c_code = t_code.cell(0, 0)
        c_code.width = Inches(6.5)
        set_cell_background(c_code, "2D3748")
        set_cell_margins(c_code, top=80, bottom=80, left=120, right=120)
        p_c = c_code.paragraphs[0]
        p_c.paragraph_format.space_after = Pt(0)
        r_c = p_c.add_run(cmd)
        r_c.font.name = 'Courier New'
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # -------------------------------------------------------------
    # SECCIÓN 6: CONCLUSIÓN Y MATRIZ DE ENTREGABLES
    # -------------------------------------------------------------
    h6 = doc.add_heading("6. Conclusiones y Cumplimiento de Entregables", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(8)

    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.line_spacing = 1.15
    p_conc.paragraph_format.space_after = Pt(6)
    p_conc.add_run(
        "La solución presentada satisface el 100% de las exigencias teóricas y prácticas planteadas en la consigna de la PFO 3. "
        "El diseño arquitectónico propuesto supera los límites del esquema monolítico previo al introducir redundancia, "
        "desacoplamiento asíncrono y escalabilidad horizontal. Por su parte, la implementación funcional en Python demuestra "
        "el correcto manejo de sockets a bajo nivel, la concurrencia de clientes y el paralelismo de ejecución en worker pools."
    )

    # Entregables summary table
    entreg_table = doc.add_table(rows=4, cols=3)
    entreg_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    entreg_table.autofit = False
    
    for idx, t in enumerate(["Ítem Requerido", "Archivo / Recurso", "Estado"]):
        c = entreg_table.rows[0].cells[idx]
        c.width = [Inches(2.0), Inches(3.2), Inches(1.3)][idx]
        set_cell_background(c, "1A365D")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(t)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    entreg_rows = [
        ("Diagrama del Sistema Distribuido", "diagrama_arquitectura.png / .svg", "✅ Cumplido"),
        ("Código Servidor y Cliente", "servidor.py y cliente.py", "✅ Cumplido"),
        ("Repositorio en GitHub", "https://github.com/guillote1977/ProgSRedes-PFO3", "✅ Cumplido")
    ]

    for row_idx, (it, f, st) in enumerate(entreg_rows, start=1):
        row = entreg_table.rows[row_idx]
        for col_idx, val in enumerate([it, f, st]):
            c = row.cells[col_idx]
            c.width = [Inches(2.0), Inches(3.2), Inches(1.3)][col_idx]
            set_cell_background(c, "F7FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            elif col_idx == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(0x27, 0x67, 0x49)
            else:
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    output_path = "PFO3_Presentacion_Sistema_Distribuido.docx"
    doc.save(output_path)
    print(f"[OK] Archivo DOCX generado exitosamente: {output_path}")

if __name__ == '__main__':
    build_presentation_docx()
